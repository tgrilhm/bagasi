"""Export kloter ke Excel — format kolom sesuai kebutuhan operasional."""
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .formatters import berat_fmt, rupiah, tanggal_indo, tanggal_waktu

COMPANY   = "Amanah Baggage"
_BLUE     = "1A3A5C"
_BLUE_LT  = "E8F0FC"
_WHITE    = "FFFFFF"
_GRAY     = "F4F7FB"
_GRAY2    = "EEF3FB"
_BORDER   = "D0E0F0"
_GREEN    = "1E8449"


def _thin(color: str = _BORDER) -> Border:
    s = Side(border_style="thin", color=color)
    return Border(left=s, right=s, top=s, bottom=s)


def export_kloter(kloter: dict, paket_list: list[dict],
                  summary: dict, output_path: Path) -> Path:
    """
    Export kloter ke Excel dengan kolom:
    ID | No Resi | Kloter | Nama (Pengirim) | No HP (Pengirim) |
    Jenis | Berat | Penerima | Bayar | Rekening Tujuan |
    Tgl Dibuat | Tgl Selesai
    """
    wb = Workbook()

    # ── Sheet 1: Info Kloter ──────────────────────────────────────────
    ws1 = wb.active
    ws1.title = "Info Kloter"
    ws1.sheet_view.showGridLines = False

    # Header perusahaan
    ws1.merge_cells("A1:F1")
    c = ws1["A1"]
    c.value     = COMPANY
    c.font      = Font(bold=True, size=16, color=_BLUE, name="Calibri")
    c.alignment = Alignment(horizontal="center", vertical="center")
    c.fill      = PatternFill("solid", fgColor=_BLUE_LT)
    ws1.row_dimensions[1].height = 30

    ws1.merge_cells("A2:F2")
    c2 = ws1["A2"]
    c2.value     = "Laporan Kloter Pengiriman"
    c2.font      = Font(size=11, color=_BLUE, italic=True, name="Calibri")
    c2.alignment = Alignment(horizontal="center")

    # Info detail kloter
    info = [
        ("Nama Kloter",    kloter["nama_kloter"]),
        ("Status",         kloter["status"]),
        ("Tanggal Dibuat", tanggal_waktu(kloter["tanggal_dibuat"])),
        ("Dibuat Oleh",    kloter["dibuat_oleh"]),
        ("Catatan",        kloter.get("catatan") or "-"),
        ("Tgl Selesai",
         tanggal_waktu(kloter["tanggal_selesai"]) if kloter.get("tanggal_selesai") else "-"),
    ]
    for i, (label, val) in enumerate(info, start=4):
        lc = ws1.cell(row=i, column=1, value=label)
        lc.font = Font(bold=True, name="Calibri")
        lc.fill = PatternFill("solid", fgColor=_GRAY)
        lc.border = _thin()
        vc = ws1.cell(row=i, column=2, value=val)
        vc.border = _thin()
        vc.font   = Font(name="Calibri")

    # Summary cards
    sum_labels = [("Total Paket", summary["total_paket"]),
                  ("Total Berat", f"{summary['total_berat']:.2f} kg"),
                  ("Total Pendapatan", rupiah(summary["total_pendapatan"]))]
    for col_off, (lbl, val) in enumerate(sum_labels, start=4):
        lc = ws1.cell(row=4, column=col_off, value=lbl)
        lc.font = Font(bold=True, name="Calibri", color=_BLUE)
        vc = ws1.cell(row=5, column=col_off, value=val)
        vc.font = Font(bold=True, size=13, name="Calibri",
                       color=_GREEN if "Pendapatan" in lbl else _BLUE)

    ws1.column_dimensions["A"].width = 20
    ws1.column_dimensions["B"].width = 30
    for col in ["D","E","F"]:
        ws1.column_dimensions[col].width = 22

    # ── Sheet 2: Daftar Paket ────────────────────────────────────────
    ws2 = wb.create_sheet("Daftar Paket")
    ws2.sheet_view.showGridLines = False

    # Header judul
    ws2.merge_cells("A1:L1")
    c1 = ws2["A1"]
    c1.value     = f"{COMPANY}  —  {kloter['nama_kloter']}"
    c1.font      = Font(bold=True, size=13, color=_WHITE, name="Calibri")
    c1.fill      = PatternFill("solid", fgColor=_BLUE)
    c1.alignment = Alignment(horizontal="center", vertical="center")
    ws2.row_dimensions[1].height = 26

    # Sub-header info
    ws2.merge_cells("A2:L2")
    c2 = ws2["A2"]
    c2.value     = (f"Status: {kloter['status']}   |   "
                    f"Dicetak: {datetime.now().strftime('%d/%m/%Y %H:%M')}   |   "
                    f"Total Paket: {summary['total_paket']}   |   "
                    f"Total Pendapatan: {rupiah(summary['total_pendapatan'])}")
    c2.font      = Font(size=9, color=_BLUE, italic=True, name="Calibri")
    c2.fill      = PatternFill("solid", fgColor=_BLUE_LT)
    c2.alignment = Alignment(horizontal="center")

    # Kolom headers sesuai permintaan:
    # ID | No Resi | Kloter | Nama | No HP | Jenis | Berat |
    # Penerima | Bayar | Rekening Tujuan | Tgl Dibuat | Tgl Selesai
    headers = [
        ("ID",              9),
        ("No Resi",        16),
        ("Kloter",         20),
        ("Nama",           22),   # Nama Pengirim
        ("No HP",          16),   # No HP Pengirim
        ("Jenis",          12),   # Kategori
        ("Berat",          10),
        ("Penerima",       22),   # Nama Penerima
        ("Bayar",          16),   # Total Harga
        ("Rekening Tujuan",22),   # Kolom baru — bisa diisi manual di Excel
        ("Tgl Dibuat",     18),
        ("Tgl Selesai",    18),   # tanggal_selesai kloter
    ]

    # Isi header row (row 3)
    for col, (h, w) in enumerate(headers, 1):
        c = ws2.cell(row=3, column=col, value=h)
        c.font      = Font(bold=True, color=_WHITE, name="Calibri", size=10)
        c.fill      = PatternFill("solid", fgColor=_BLUE)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border    = _thin(_BLUE)
        ws2.column_dimensions[get_column_letter(col)].width = w
    ws2.row_dimensions[3].height = 22

    # Data rows (mulai row 4)
    tgl_selesai_kloter = (tanggal_waktu(kloter["tanggal_selesai"])
                          if kloter.get("tanggal_selesai") else "-")

    for idx, p in enumerate(paket_list, 1):
        r   = idx + 3
        alt = PatternFill("solid", fgColor=_GRAY2) if idx % 2 == 0 else None

        values = [
            p["id"],
            p["no_resi"],
            kloter["nama_kloter"],
            p["nama_pengirim"],
            p["no_hp_pengirim"],
            p["kategori"],
            p["berat_kg"],
            p["nama_penerima"],
            p["total_harga"],
            "",                             # Rekening Tujuan — diisi manual
            tanggal_waktu(p["tanggal_dibuat"]),
            tgl_selesai_kloter,
        ]
        for col, val in enumerate(values, 1):
            cell = ws2.cell(row=r, column=col, value=val)
            cell.border = _thin()
            cell.font   = Font(name="Calibri", size=9)
            if alt:
                cell.fill = alt
            # Format angka
            if col == 7:  # Berat
                cell.number_format = "0.00"
                cell.alignment = Alignment(horizontal="center")
            elif col == 9:  # Bayar
                cell.number_format = "#,##0"
                cell.alignment = Alignment(horizontal="right")
                cell.font = Font(name="Calibri", size=9, bold=True)
            elif col in (1, 2, 10):  # ID, No Resi, Rekening Tujuan
                cell.alignment = Alignment(horizontal="center")
            elif col == 6:  # Jenis
                cell.alignment = Alignment(horizontal="center")

    # Baris total
    total_row = len(paket_list) + 4
    ws2.cell(row=total_row, column=8, value="TOTAL").font = Font(
        bold=True, name="Calibri", color=_BLUE)
    tc = ws2.cell(row=total_row, column=9, value=summary["total_pendapatan"])
    tc.font          = Font(bold=True, name="Calibri", color=_GREEN, size=11)
    tc.number_format = "#,##0"
    tc.alignment     = Alignment(horizontal="right")

    # Timestamp footer
    ws2.cell(
        row=total_row + 2, column=1,
        value=f"Dicetak: {datetime.now().strftime('%d/%m/%Y %H:%M')}  |  {COMPANY}"
    ).font = Font(italic=True, color="888888", name="Calibri", size=8)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(output_path))
    return output_path
